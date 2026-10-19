"""Dataset batches: find meter photos and labels in folders/ZIPs, compare readings, export CSV and Excel."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path, PurePosixPath
from statistics import mean

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
# The Kaggle "Water Meters" dataset keeps masks and collages next to the photos; they are not meter photos.
SKIP_DIRS = {"mask", "masks", "collage", "collages", "__macosx"}
NAME_COLUMNS = ("photo_name", "image_name", "filename", "file_name", "file", "image", "name")
VALUE_COLUMNS = ("value", "reading", "expected", "label", "meter_value")
VALUE_IN_NAME = re.compile(r"value_(\d+)(?:_(\d+))?", re.IGNORECASE)
MAX_IMAGE_BYTES = 12 * 1024 * 1024
# On unseen Kaggle test photos, readings under 45% confidence were right 27% of the time, 45%+ about 94%.
LOW_CONFIDENCE = 45.0

# (key, header, Excel column width). CSV uses the keys as headers, Excel uses the friendly headers.
COLUMNS = [
    ("id", "ID", 8), ("batch_id", "Batch", 8), ("meter_id", "Meter ID", 16), ("source_name", "Source file", 34),
    ("reading", "Reading", 14), ("value", "Value", 12), ("confidence", "Confidence %", 13),
    ("expected", "Expected", 12), ("match", "Match", 9), ("difference", "Difference", 12), ("split", "Model split", 12),
    ("corrected", "Corrected", 10), ("roi_source", "ROI", 11), ("note", "Note", 30),
    ("captured_at", "Captured at", 20), ("image_name", "Stored image", 30),
]
CLI_COLUMNS = [c for c in COLUMNS if c[0] not in {"id", "batch_id", "corrected", "image_name"}]


def _natural_key(name: str):
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", name)]


def _clean_parts(name: str) -> list[str]:
    return [part for part in name.replace("\\", "/").split("/") if part not in {"", ".", ".."}]


def is_image_name(name: str) -> bool:
    parts = _clean_parts(name)
    if not parts or parts[-1].startswith("."):
        return False
    return PurePosixPath(parts[-1]).suffix.lower() in IMAGE_SUFFIXES and not {p.lower() for p in parts[:-1]} & SKIP_DIRS


def parse_labels(text: str) -> dict[str, str]:
    """Read a labels CSV (for example Kaggle data.csv: photo_name,value) into {file name: expected value}."""
    rows = csv.DictReader(io.StringIO(text.lstrip("\ufeff")))
    fields = {(f or "").strip().lower(): f for f in rows.fieldnames or []}
    name_col = next((fields[c] for c in NAME_COLUMNS if c in fields), None)
    value_col = next((fields[c] for c in VALUE_COLUMNS if c in fields), None)
    if not name_col or not value_col:
        return {}
    labels = {}
    for row in rows:
        name, value = (row.get(name_col) or "").strip(), clean_number(row.get(value_col))
        if name and value:
            labels[_clean_parts(name)[-1]] = value
    return labels


def clean_number(value) -> str:
    """'305.16200000000003' -> '305.162'; keeps the label's own precision and drops float noise."""
    try:
        number = float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return ""
    return f"{round(number, 6):.6f}".rstrip("0").rstrip(".")


def expected_for(name: str, labels: dict[str, str]) -> str:
    base = _clean_parts(name)[-1] if _clean_parts(name) else name
    if base in labels:
        return labels[base]
    match = VALUE_IN_NAME.search(base)  # Kaggle file names carry the value: id_53_value_595_825.jpg
    return clean_number(f"{match.group(1)}.{match.group(2) or 0}") if match else ""


def to_number(reading) -> float | None:
    text = str(reading or "").strip()
    return float(text) if re.fullmatch(r"\d+(\.\d+)?", text) else None


class Dataset:
    """Meter photos and label CSVs inside a folder or ZIP. A folder may itself contain ZIPs."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.labels: dict[str, str] = {}
        self._entries: list[tuple[str, object]] = []
        self._zips: list[zipfile.ZipFile] = []
        if self.path.is_dir():
            for file in sorted(self.path.rglob("*"), key=lambda p: _natural_key(str(p))):
                if file.is_file():
                    self._add_file(file, file.relative_to(self.path).as_posix())
        elif self.path.is_file():
            self._add_file(self.path, self.path.name)
        else:
            raise FileNotFoundError(f"{self.path} does not exist")
        self._entries.sort(key=lambda e: _natural_key(e[0]))

    def _add_file(self, file: Path, name: str):
        suffix = file.suffix.lower()
        if suffix == ".zip":
            archive = zipfile.ZipFile(file)
            self._zips.append(archive)
            for info in archive.infolist():
                if info.is_dir():
                    continue
                if info.filename.lower().endswith(".csv") and not {p.lower() for p in _clean_parts(info.filename)[:-1]} & SKIP_DIRS:
                    self.labels.update(parse_labels(archive.read(info).decode("utf-8", errors="replace")))
                elif is_image_name(info.filename):
                    self._entries.append((info.filename, (archive, info)))
        elif suffix == ".csv":
            self.labels.update(parse_labels(file.read_text(encoding="utf-8", errors="replace")))
        elif is_image_name(name):
            self._entries.append((name, file))

    @property
    def names(self) -> list[str]:
        return [name for name, _ in self._entries]

    def __len__(self):
        return len(self._entries)

    def __iter__(self):
        """Yields (name, bytes loader). Loading is lazy so a large ZIP is never held in memory."""
        for name, source in self._entries:
            yield name, (lambda s=source: self._load(s))

    @staticmethod
    def _load(source) -> bytes:
        if isinstance(source, Path):
            if source.stat().st_size > MAX_IMAGE_BYTES:
                raise ValueError("Image is larger than 12 MB")
            return source.read_bytes()
        archive, info = source
        if info.file_size > MAX_IMAGE_BYTES:
            raise ValueError("Image is larger than 12 MB")
        with archive.open(info) as handle:
            return handle.read(MAX_IMAGE_BYTES + 1)

    def close(self):
        for archive in self._zips:
            archive.close()

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()


def load_split(models_dir: Path) -> dict[str, str]:
    """{file name: train/val/test} written by the training script, so exports can flag images the model learned from."""
    path = Path(models_dir) / "meter_split.json"
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return {name: split for split, names in data.items() if isinstance(names, list) for name in names}


def enrich(row: dict, split: dict[str, str] | None = None) -> dict:
    """Adds numeric value, match and difference against the expected reading."""
    row = dict(row)
    value, expected = to_number(row.get("reading")), to_number(row.get("expected"))
    row["value"] = value
    row["difference"] = round(value - expected, 6) if value is not None and expected is not None else None
    row["match"] = "" if expected is None else ("Yes" if row["difference"] is not None and abs(row["difference"]) < 1e-6 else "No")
    if split:
        parts = _clean_parts(row.get("source_name") or "")
        row["split"] = split.get(parts[-1], "") if parts else ""
    return row


def summarize(rows: list[dict]) -> list[tuple[str, object]]:
    total = len(rows)
    read = sum(1 for r in rows if r.get("reading"))
    items: list[tuple[str, object]] = [("Images", total), ("Readings extracted", read), ("No reading", total - read)]
    confidences = [float(r["confidence"]) for r in rows if r.get("confidence") is not None]
    if confidences:
        items.append(("Average confidence %", round(mean(confidences), 1)))
        items.append((f"Low confidence (under {LOW_CONFIDENCE:g}%) — check these", sum(c < LOW_CONFIDENCE for c in confidences)))
    failed = sum(1 for r in rows if r.get("note"))
    if failed:
        items.append(("Images that could not be read", failed))
    corrected = sum(1 for r in rows if r.get("corrected"))
    if corrected:
        items.append(("Manually corrected", corrected))
    labelled = [r for r in rows if r.get("match")]
    if labelled:
        matched = sum(1 for r in labelled if r["match"] == "Yes")
        items += [("Images with expected value", len(labelled)), ("Exact matches", matched),
                  ("Accuracy %", round(100 * matched / len(labelled), 1))]
        for split, label in (("test", "never seen in training"), ("val", "validation"), ("train", "used for training")):
            subset = [r for r in labelled if r.get("split") == split]
            if subset:
                hits = sum(1 for r in subset if r["match"] == "Yes")
                items.append((f"Accuracy % — {split} images ({label})", round(100 * hits / len(subset), 1)))
    return items


def active_columns(columns, rows):
    return [c for c in columns if c[0] != "split" or any(r.get("split") for r in rows)]


def _safe_text(value):
    # Stops spreadsheet apps from treating a file name like "=HYPERLINK(...)" as a formula.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def csv_text(rows: list[dict], columns=COLUMNS) -> str:
    columns = active_columns(columns, rows)
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([key for key, _, _ in columns])
    for row in rows:
        writer.writerow(["" if row.get(key) is None else _safe_text(row.get(key)) for key, _, _ in columns])
    return out.getvalue()


def xlsx_bytes(rows: list[dict], columns=COLUMNS, summary=None, title="Water meter readings") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = active_columns(columns, rows)
    header_fill, header_font = PatternFill("solid", fgColor="087F79"), Font(bold=True, color="FFFFFF")
    yes_fill, no_fill = PatternFill("solid", fgColor="DCF7EE"), PatternFill("solid", fgColor="FFE3E1")
    book = Workbook()
    sheet = book.active
    sheet.title = "Readings"
    sheet.append([header for _, header, _ in columns])
    for index, (key, _, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill, cell.font, cell.alignment = header_fill, header_font, Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    formats = {"confidence": "0.0", "value": "0.000", "difference": "0.000"}
    for row_index, row in enumerate(rows, start=2):
        for col_index, (key, _, _) in enumerate(columns, start=1):
            value = row.get(key)
            if key == "corrected":
                value = "Yes" if value else ""
            cell = sheet.cell(row=row_index, column=col_index)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # never evaluate as a formula
            if key in formats and isinstance(value, (int, float)):
                cell.number_format = formats[key]
            if key == "match" and value:
                cell.fill = yes_fill if value == "Yes" else no_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    info = book.create_sheet("Summary")
    info.append([title])
    info["A1"].font = Font(bold=True, size=14)
    info.append([])
    for label, value in summary if summary is not None else summarize(rows):
        info.append([label, value])
    info.column_dimensions["A"].width = 46
    info.column_dimensions["B"].width = 16
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
