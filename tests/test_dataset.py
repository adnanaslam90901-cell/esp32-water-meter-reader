import csv
import io
import zipfile

import cv2
import numpy as np
from openpyxl import load_workbook

from app import dataset as ds


def png(code=10):
    return cv2.imencode(".png", np.full((8, 8, 3), code, np.uint8))[1].tobytes()


def test_kaggle_labels_drop_float_noise():
    text = 'photo_name,value,location\nid_851_value_305_162.jpg,305.16200000000003,"{\'type\': \'polygon\'}"\nid_407_value_21_86.jpg,21.86,x\n'
    assert ds.parse_labels(text) == {"id_851_value_305_162.jpg": "305.162", "id_407_value_21_86.jpg": "21.86"}


def test_labels_need_name_and_value_columns():
    assert ds.parse_labels("a,b\n1,2\n") == {}
    assert ds.parse_labels("﻿Filename,Reading\nfolder/x.jpg,00012.5\n") == {"x.jpg": "12.5"}


def test_expected_value_from_labels_or_kaggle_file_name():
    assert ds.expected_for("images/id_1_value_595_825.jpg", {}) == "595.825"
    assert ds.expected_for("images/id_1_value_195_1.jpg", {}) == "195.1"
    assert ds.expected_for("a/b/photo.jpg", {"photo.jpg": "7.5"}) == "7.5"
    assert ds.expected_for("photo.jpg", {}) == ""


def test_enrich_marks_match_and_difference():
    assert ds.enrich({"reading": "00021.860", "expected": "21.86"})["match"] == "Yes"
    wrong = ds.enrich({"reading": "00021.870", "expected": "21.86"})
    assert (wrong["match"], wrong["difference"], wrong["value"]) == ("No", 0.01, 21.87)
    assert ds.enrich({"reading": "", "expected": "21.86"})["match"] == "No"
    assert ds.enrich({"reading": "00021.860", "expected": ""})["match"] == ""
    assert ds.enrich({"reading": "1", "source_name": "x/id_9.jpg"}, {"id_9.jpg": "test"})["split"] == "test"


def test_dataset_reads_folders_zips_and_skips_masks(tmp_path):
    (tmp_path / "images").mkdir()
    (tmp_path / "masks").mkdir()
    (tmp_path / "images" / "id_2_value_1_5.png").write_bytes(png())
    (tmp_path / "images" / "id_10_value_3_0.png").write_bytes(png())
    (tmp_path / "masks" / "id_2_value_1_5.png").write_bytes(png())
    (tmp_path / "notes.txt").write_text("not an image")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("more/photo.jpg", png())
        archive.writestr("more/collage/photo.jpg", png())
        archive.writestr("more/labels.csv", "file,value\nphoto.jpg,42.125\n")
    (tmp_path / "extra.zip").write_bytes(buffer.getvalue())

    with ds.Dataset(tmp_path) as dataset:
        assert dataset.names == ["images/id_2_value_1_5.png", "images/id_10_value_3_0.png", "more/photo.jpg"]
        assert dataset.labels == {"photo.jpg": "42.125"}
        assert all(len(load()) > 0 for _, load in dataset)


def test_exports_guard_formulas_and_include_summary():
    rows = [ds.enrich({"id": 1, "source_name": "=HYPERLINK(\"x\")", "reading": "00001.500", "expected": "1.5", "confidence": 88.0}),
            ds.enrich({"id": 2, "source_name": "b.jpg", "reading": "00002.000", "expected": "1.5", "confidence": 40.0})]
    parsed = list(csv.DictReader(io.StringIO(ds.csv_text(rows))))
    assert parsed[0]["source_name"].startswith("'=")
    assert [r["match"] for r in parsed] == ["Yes", "No"]
    assert "split" not in parsed[0]

    book = load_workbook(io.BytesIO(ds.xlsx_bytes(rows)))
    sheet = book["Readings"]
    headers = [c.value for c in sheet[1]]
    assert headers[:5] == ["ID", "Batch", "Meter ID", "Source file", "Reading"]
    assert sheet.cell(row=2, column=4).data_type == "s"
    summary = {r[0].value: r[1].value for r in book["Summary"].iter_rows(min_row=3)}
    assert summary["Accuracy %"] == 50.0 and summary["Images"] == 2
