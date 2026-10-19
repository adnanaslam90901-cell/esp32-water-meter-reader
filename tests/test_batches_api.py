import csv
import io
import time
import zipfile

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main

FAKE_READINGS = {10: "00012.345", 20: "00007.400", 30: ""}


def png(code):
    return cv2.imencode(".png", np.full((8, 8, 3), code, np.uint8))[1].tobytes()


@pytest.fixture
def client(tmp_path, monkeypatch):
    (tmp_path / "captures").mkdir()
    monkeypatch.setattr(main, "DB", tmp_path / "readings.db")
    monkeypatch.setattr(main, "CAPTURES", tmp_path / "captures")
    monkeypatch.setattr(main, "BATCH_UPLOADS", tmp_path / "uploads")
    monkeypatch.setattr(main, "MODELS", tmp_path / "models")

    def fake_read(image):
        reading = FAKE_READINGS[int(image[0, 0, 0])]
        return {"reading": reading, "confidence": 90.0, "variant": "fake", "raw": reading,
                "roi": {"x": 0, "y": 0, "w": 8, "h": 8, "source": "model"}, "image_width": 8, "image_height": 8}

    monkeypatch.setattr(main.reader, "read", fake_read)
    with TestClient(main.app) as test_client:
        yield test_client


def wait_until_finished(client, batch_id):
    for _ in range(100):
        batch = client.get(f"/api/batches/{batch_id}").json()
        if batch["status"] not in ("queued", "running"):
            return batch
        time.sleep(0.05)
    raise AssertionError("batch did not finish")


def kaggle_zip():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("WaterMeters/data.csv", "photo_name,value,location\nid_2_value_7_5.png,7.5,x\n")
        archive.writestr("WaterMeters/images/id_1_value_12_345.png", png(10))  # expected from file name, matches
        archive.writestr("WaterMeters/images/id_2_value_7_5.png", png(20))     # expected from data.csv, wrong reading
        archive.writestr("WaterMeters/images/broken.jpg", b"not really a jpeg")
        archive.writestr("WaterMeters/images/unlabelled.png", png(30))
        archive.writestr("WaterMeters/masks/id_1_value_12_345.png", png(10))  # masks are not meter photos
    return buffer.getvalue()


def test_zip_batch_reads_compares_and_exports(client):
    created = client.post("/api/batches?meter_id=kaggle", files=[("files", ("WaterMeters.zip", kaggle_zip(), "application/zip"))])
    assert created.status_code == 200, created.text
    assert (created.json()["total"], created.json()["labelled"]) == (4, 2)

    batch = wait_until_finished(client, created.json()["id"])
    assert (batch["status"], batch["processed"], batch["failed"]) == ("done", 4, 1)
    summary = {s["label"]: s["value"] for s in batch["summary"]}
    assert summary["Accuracy %"] == 50.0 and summary["Exact matches"] == 1

    mismatches = client.get(f"/api/batches/{batch['id']}/readings?show=mismatch").json()
    assert [r["source_name"] for r in mismatches] == ["WaterMeters/images/id_2_value_7_5.png"]

    exported = list(csv.DictReader(io.StringIO(client.get(f"/api/export.csv?batch_id={batch['id']}").text)))
    by_file = {r["source_name"].split("/")[-1]: r for r in exported}
    assert by_file["id_1_value_12_345.png"]["match"] == "Yes"
    assert by_file["id_2_value_7_5.png"]["match"] == "No"
    assert by_file["broken.jpg"]["note"] and by_file["broken.jpg"]["reading"] == ""
    assert by_file["unlabelled.png"]["expected"] == ""

    response = client.get(f"/api/export.xlsx?batch_id={batch['id']}")
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    book = load_workbook(io.BytesIO(response.content))
    assert book["Readings"].max_row == 5 and "Summary" in book.sheetnames

    assert client.get("/api/readings").json() == []  # batch rows stay out of the live capture history
    assert client.get("/api/batches").json()[0]["accuracy"] == 50.0
    assert client.delete(f"/api/batches/{batch['id']}").status_code == 200
    assert client.get("/api/batches").json() == []
    assert not list(main.CAPTURES.iterdir())


def test_folder_upload_keeps_paths_but_not_traversal(client):
    files = [("files", ("../../escape/id_1_value_12_345.png", png(10), "image/png")),
             ("files", ("meters/labels.csv", b"filename,reading\nid_1_value_12_345.png,99\n", "text/csv"))]
    created = client.post("/api/batches", files=files).json()
    batch = wait_until_finished(client, created["id"])
    rows = client.get(f"/api/batches/{batch['id']}/readings").json()
    assert rows[0]["source_name"] == "escape/id_1_value_12_345.png"
    assert rows[0]["expected"] == "99" and rows[0]["match"] == "No"  # labels CSV wins over the file name


def test_upload_without_photos_is_rejected(client):
    response = client.post("/api/batches", files=[("files", ("notes.csv", b"a,b\n", "text/csv"))])
    assert response.status_code == 400
